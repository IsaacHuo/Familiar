"""Render a sourced research document. Content comes from the task, never a fixture."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    payload = json.loads(Path(args.input).read_text(encoding="utf-8"))
    assert payload.get("title") and payload.get("sections"), "Title and sections are required"
    document = Document()
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.font.name = "Arial"
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), "PingFang SC")
    document.add_heading(payload["title"], 0)
    for section in payload["sections"]:
        assert section.get("heading") and section.get("paragraphs"), "Empty section"
        document.add_heading(section["heading"], 1)
        for paragraph in section["paragraphs"]:
            assert isinstance(paragraph, str) and paragraph.strip(), "Empty paragraph"
            document.add_paragraph(paragraph)
    if payload.get("sources"):
        document.add_heading("参考资料", 1)
        for source in payload["sources"]:
            assert source["url"].startswith("https://"), "Source URL must be HTTPS"
            document.add_paragraph(source["title"] + "\n" + source["url"])
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    reopened = Document(output)
    assert reopened.paragraphs and output.stat().st_size > 0
    print(json.dumps({"path": str(output), "bytes": output.stat().st_size, "paragraphs": len(reopened.paragraphs)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
